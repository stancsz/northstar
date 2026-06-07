#!/usr/bin/env python3
"""
Layer 1 / UNDERSTAND — Standards Conformance mode ingestion.

Extracts text from a standards document (PDF, DOCX, MD, HTML, TXT), parses
clause headings (e.g. "5.1.1 Functional completeness"), and finds normative
sentences (Chinese 应/应按/不得/必须/应符合 + English shall/must). Writes the
same `_prd_text.txt` / `_prd_candidates.json` files as ingest_prd.py so
downstream tooling stays single-path, plus a `standards_index.json` at the
run root with metadata used by verify.py and make_report.py.

The script never invents semantics. Claude reads the candidates and authors
the requirement graph with `type: "conformance_criterion"`, `standard_id`,
and `clause` populated.

Usage:
    python ingest_standard.py <standard-path> --out <run-dir> [--standard-id "GB/T 25000.51-2016"]
"""
import argparse, json, os, re, sys
from _common import info, die, save, now

# ---------- text extraction (mirror of ingest_prd.extract) ----------
def extract(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            die("pypdf not installed. Run: pip install pypdf --break-system-packages")
        return "\n".join((p.extract_text() or "") for p in PdfReader(path).pages)
    if ext == ".docx":
        try:
            import docx
        except ImportError:
            die("python-docx not installed. Run: pip install python-docx --break-system-packages")
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if ext in (".html", ".htm"):
        try:
            from bs4 import BeautifulSoup
            with open(path, encoding="utf-8") as f:
                return BeautifulSoup(f.read(), "html.parser").get_text("\n")
        except ImportError:
            pass
    with open(path, encoding="utf-8", errors="replace") as f:
        return f.read()

# ---------- standards-specific parsing ----------
# Clause heading: 1 to 4 dot-separated numeric parts followed by a space and a non-empty title.
# Anchored to start of line so we don't pick up "1.2.3" in prose; allows leading whitespace.
CLAUSE_RE = re.compile(r"^\s*(\d+(?:\.\d+){0,3})\s+(\S.+?)\s*$")
# Normative modal verbs (Chinese + English).
MODAL_RE = re.compile(
    r"应(?!\w)|应按|不得|必须|应符合|shall|must", re.IGNORECASE
)
# Standard-ID pattern for filename / first-page sniff.
STANDARD_ID_RE = re.compile(
    r"\b(GB/T|GB|IEC|ISO|ISO/IEC|EN|BS|JIS|DIN|GOST|KS|NF|UNE|NBR|NBN|SS|AS/NZS)\s*"
    r"[A-Z]?[\d.]+(?:[-:—–]\d+)?(?:[-:—–]\d+)?\b",
    re.IGNORECASE,
)

def parse_clauses(text):
    """Return [(clause_id, title, line_index), ...] in document order."""
    out = []
    for i, line in enumerate(text.splitlines()):
        m = CLAUSE_RE.match(line)
        if m:
            out.append((m.group(1), m.group(2).strip(), i))
    return out

def extract_normative_sentences(text, clauses):
    """Split text on sentence-ish boundaries and keep sentences that contain a
    normative modal. Attaches the *current* clause at the time of the sentence
    so the agent can author requirements with the right `clause` field.

    Returns: [{clause, text, line}, ...]
    """
    # Walk line-by-line, tracking the most recent clause heading.
    clause_at = [None] * len(text.splitlines())
    cur = None
    for idx, line in enumerate(text.splitlines()):
        m = CLAUSE_RE.match(line)
        if m:
            cur = m.group(1)
        clause_at[idx] = cur

    out, seen = [], set()
    # Split on Chinese & English sentence boundaries and newlines.
    chunks = re.split(r"(?<=[。！？!?])\s+|(?<=[.!?])\s+|\n+", text)
    for chunk in chunks:
        s = chunk.strip().replace("\n", " ")
        s = re.sub(r"\s+", " ", s)
        if 6 <= len(s) <= 400 and MODAL_RE.search(s) and s.lower() not in seen:
            seen.add(s.lower())
            # Approximate the source line for the agent.
            approx_line = text[: text.find(s)].count("\n") + 1 if s in text else None
            clause = clause_at[approx_line - 1] if approx_line else None
            out.append({"clause": clause, "text": s, "line": approx_line})
    return out

def guess_standard_id(path, text):
    """Try filename, then first 4k chars of text."""
    name = os.path.basename(path)
    # Filename: replace separators commonly used in standard IDs with spaces.
    cleaned = re.sub(r"[_\-.]+", " ", os.path.splitext(name)[0])
    m = STANDARD_ID_RE.search(cleaned.upper())
    if m:
        return m.group(0).strip()
    head = text[:4000]
    m = STANDARD_ID_RE.search(head)
    return m.group(0).strip() if m else None

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("standard", help="path to standards document (PDF/DOCX/MD/HTML/TXT)")
    ap.add_argument("--out", required=True, help="run directory")
    ap.add_argument("--standard-id", default=None,
                    help="override auto-detected standard ID (e.g. 'GB/T 25000.51-2016')")
    a = ap.parse_args()
    if not os.path.exists(a.standard):
        die(f"Standard not found: {a.standard}")
    os.makedirs(a.out, exist_ok=True)
    text = extract(a.standard)
    if not text.strip():
        die("No text extracted. If this is a scanned PDF, OCR it first (see pdf skill).")

    with open(os.path.join(a.out, "_prd_text.txt"), "w", encoding="utf-8") as f:
        f.write(text)

    clauses = parse_clauses(text)
    sentences = extract_normative_sentences(text, clauses)
    standard_id = a.standard_id or guess_standard_id(a.standard, text)

    # Reuse the same candidates filename for downstream simplicity.
    save({
        "source": a.standard,
        "generated_at": now(),
        "char_count": len(text),
        "kind": "standard",
        "standard_id": standard_id,
        "clauses": [{"id": c[0], "title": c[1], "line": c[2]} for c in clauses],
        "candidates": sentences,  # normative sentences (one requirement seed each)
    }, os.path.join(a.out, "_prd_candidates.json"))

    save({
        "standard_id": standard_id,
        "standard_name": None,  # agent fills this in 01_requirements.json's `standard_name` per req
        "edition": None,
        "publisher": None,
        "source_path": a.standard,
        "clause_count": len(clauses),
        "normative_sentence_count": len(sentences),
        "generated_at": now(),
    }, os.path.join(a.out, "standards_index.json"))

    info(f"Extracted {len(text)} chars; {len(clauses)} clause headings; "
         f"{len(sentences)} normative sentences.")
    if standard_id:
        info(f"Detected standard ID: {standard_id}")
    info("Next: read _prd_text.txt + _prd_candidates.json and author 01_requirements.json "
         "with type='conformance_criterion', standard_id, and clause populated. "
         "See references/standards_to_testcases.md and references/schemas.md.")

if __name__ == "__main__":
    main()
