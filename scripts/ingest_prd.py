#!/usr/bin/env python3
"""
Layer 1 / UNDERSTAND — PRD ingestion.

Extracts text from a PRD (PDF, DOCX, Markdown, TXT, HTML/Confluence/Notion export)
and emits a plain-text dump plus a heuristic first pass of candidate requirement
sentences. Claude reads these and authors the real requirement graph
(01_requirements.json) — this script just removes the mechanical parsing burden
and never invents semantics.

Usage:
    python ingest_prd.py <prd-path> --out <run-dir>

Writes:
    <run-dir>/_prd_text.txt        full extracted text
    <run-dir>/_prd_candidates.json heuristic candidate requirements
"""
import argparse, json, os, re, sys
from _common import info, die, save, now

MODAL = re.compile(r"\b(must|shall|should|will|needs? to|is required to|the system|"
                   r"the user can|the agent|when .+ then)\b", re.IGNORECASE)

def extract(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            die("pypdf not installed. Run: pip install pypdf --break-system-packages")
        return "\n".join((p.extract_text() or "") for p in PdfReader(path).pages)
    if ext == ".docx":
        try:
            import docx
        except ImportError:
            die("python-docx not installed. Run: pip install python-docx --break-system-packages")
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if ext in (".html", ".htm"):
        try:
            from bs4 import BeautifulSoup
            with open(path, encoding="utf-8") as f:
                return BeautifulSoup(f.read(), "html.parser").get_text("\n")
        except ImportError:
            pass  # fall through to raw read
    # md, txt, json export, or html fallback
    with open(path, encoding="utf-8", errors="replace") as f:
        return f.read()

def candidates(text):
    out, seen = [], set()
    # split on sentence-ish and bullet boundaries
    for chunk in re.split(r"(?:\n[-*•\d.\)]+\s+)|(?<=[.!?])\s+|\n{2,}", text):
        s = chunk.strip().replace("\n", " ")
        if 8 <= len(s) <= 400 and MODAL.search(s) and s.lower() not in seen:
            seen.add(s.lower())
            out.append(s)
    return out

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("prd")
    ap.add_argument("--out", required=True, help="run directory")
    a = ap.parse_args()
    if not os.path.exists(a.prd):
        die(f"PRD not found: {a.prd}")
    os.makedirs(a.out, exist_ok=True)
    text = extract(a.prd)
    if not text.strip():
        die("No text extracted. If this is a scanned PDF, OCR it first (see pdf skill).")
    with open(os.path.join(a.out, "_prd_text.txt"), "w", encoding="utf-8") as f:
        f.write(text)
    cands = candidates(text)
    save({"source": a.prd, "generated_at": now(),
          "char_count": len(text), "candidates": cands},
         os.path.join(a.out, "_prd_candidates.json"))
    info(f"Extracted {len(text)} chars; {len(cands)} candidate requirements.")
    info("Next: read _prd_text.txt and author 01_requirements.json (see references/schemas.md).")

if __name__ == "__main__":
    main()
